"""
Основные callback функции для дашборда
"""
import logging
from functools import lru_cache

import dash
from dash import Input, Output, State, callback, html, dcc
import pandas as pd
from datetime import datetime, timedelta
import json
import random

logger = logging.getLogger(__name__)

from data.queries_mssql import (
    get_orders_timely, get_avg_operation_time, get_total_earnings, get_order_accuracy,
    get_avg_productivity, get_performance_data, get_shift_comparison,
    get_problematic_hours, get_fines_data,
    get_employees_on_shift, get_employees_on_shift_new, refresh_data, get_error_hours_top_data,
    get_storage_cells_stats, get_all_storage_data, get_revision_stats, get_placement_errors,
    filter_storage_data, get_group_load_monitor, check_data_availability
)
from components.charts import (
    create_order_accuracy_chart, create_problematic_hours_chart,
    create_timeliness_chart, create_operations_type_chart,
    create_time_distribution_pie_echarts, create_idle_intervals_bar_echarts,
    create_fines_pie_chart, create_fines_amount_bar_chart,
    create_employee_fines_chart,
    create_empty_pie_chart, create_types_pie_chart, create_types_bar_chart
)
from components.tables import create_performance_table

# Импорты для клиентского кэширования и ETL статуса
from utils.etl_status import ETLChecker
from utils.client_cache import ClientCacheManager

# ============================================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ТАБЛИЦЫ СОТРУДНИКОВ
# ============================================================================
ACTIVE_STATUSES = ['На смене', 'Вышел', 'Работает']

def _calc_idle_seconds(last_op_str):
    """Рассчитывает секунды простоя от времени последней операции до сейчас"""
    if not last_op_str or last_op_str in ['--:--', 'None', '']:
        return 0
    try:
        now = datetime.now()
        parts = str(last_op_str).split(':')
        h, m = int(parts[0]), int(parts[1])
        last_op_dt = now.replace(hour=h, minute=m, second=0, microsecond=0)
        if last_op_dt > now:
            last_op_dt -= timedelta(days=1)
        return max(0, int((now - last_op_dt).total_seconds()))
    except:
        return 0

def _apply_sorting(clicked_col, current_state, pf, bf):
    """Единая логика сортировки с принудительным сбросом и числовым сравнением"""
    from data.queries_mssql import get_employees_on_shift_new

    cur_col = current_state.get('column') if current_state else None
    cur_dir = current_state.get('direction') if current_state else None
    force_asc = current_state.get('force_dir') == 'asc'  # Флаг принудительного сброса

    # Логика определения направления
    if force_asc:
        new_col, new_dir = clicked_col, 'asc'
    elif clicked_col == cur_col:
        if cur_dir == 'asc': new_dir = 'desc'
        elif cur_dir == 'desc': new_dir = 'reset'
        else: new_dir = 'asc'
    else:
        new_col, new_dir = clicked_col, 'asc'

    if new_dir == 'reset':
        new_col, new_dir = 'time', 'asc'
    else:
        new_col = clicked_col

    employees, _ = get_employees_on_shift_new()
    filtered = [e for e in employees if (not pf or pf=='all' or e.get('Должность')==pf) and (not bf or bf=='all' or e.get('Бригада')==bf)]

    # ✅ Числовой парсер времени (H:M -> минуты)
    def parse_time_to_minutes(time_str):
        if not time_str or time_str == '--:--': return 99999
        try: 
            parts = str(time_str).split(':')
            return int(parts[0]) * 60 + int(parts[1])
        except: return 99999

    if new_col == 'time':
        if new_dir == 'asc':
            # От раннего к позднему
            filtered.sort(key=lambda e: parse_time_to_minutes(e.get('Время_первой_операции', '--:--')))
        else:
            # От позднего к раннему, "--:--" и "Не вышел" строго в конце
            valid = [e for e in filtered if e.get('Статус') not in ['Не вышел'] and e.get('Время_первой_операции', '--:--') != '--:--']
            invalid = [e for e in filtered if e.get('Статус') in ['Не вышел'] or e.get('Время_первой_операции', '--:--') == '--:--']
            valid.sort(key=lambda e: parse_time_to_minutes(e.get('Время_первой_операции', '--:--')), reverse=True)
            filtered = valid + invalid

    elif new_col == 'status':
        filtered.sort(key=lambda e: 0 if e.get('Статус') in ['На смене', 'Вышел', 'Работает'] else 1, reverse=(new_dir=='desc'))

    elif new_col == 'last_time':
        if new_dir == 'asc':
            filtered.sort(key=lambda e: parse_time_to_minutes(e.get('Время_последней_операции', '--:--')))
        else:
            valid = [e for e in filtered if e.get('Время_последней_операции', '--:--') != '--:--']
            invalid = [e for e in filtered if e.get('Время_последней_операции', '--:--') == '--:--']
            valid.sort(key=lambda e: parse_time_to_minutes(e.get('Время_последней_операции', '--:--')), reverse=True)
            filtered = valid + invalid

    elif new_col == 'idle':
        def idle_key(e):
            if e.get('Статус') == 'Не вышел': return 999999
            return _calc_idle_seconds(e.get('Время_последней_операции', '--:--'))
        valid = [e for e in filtered if e.get('Статус') != 'Не вышел']
        invalid = [e for e in filtered if e.get('Статус') == 'Не вышел']
        valid.sort(key=idle_key, reverse=(new_dir=='desc'))
        filtered = valid + invalid

    # Генерация строк (без изменений)
    now_ms = int(datetime.now().timestamp() * 1000)
    rows = []
    for employee in filtered:
        status = employee.get('Статус', 'Не вышел')
        is_on_shift = status in ['На смене', 'Вышел', 'Работает']
        status_color = '#4CAF50' if is_on_shift else '#F44336'
        last_op = employee.get('Время_последней_операции', '--:--')
        idle_sec = _calc_idle_seconds(last_op) if is_on_shift else 0

        rows.append(html.Tr([
            html.Td(employee.get('ФИО', ''), style={'padding': '8px', 'borderBottom': '1px solid #eee'}),
            html.Td(employee.get('Должность', ''), style={'padding': '8px', 'borderBottom': '1px solid #eee'}),
            html.Td(employee.get('Бригада', ''), style={'padding': '8px', 'borderBottom': '1px solid #eee'}),
            html.Td(status, style={'padding': '8px', 'borderBottom': '1px solid #eee', 'color': status_color, 'fontWeight': 'bold'}),
            html.Td(employee.get('Время_первой_операции', '--:--'), style={'padding': '8px', 'borderBottom': '1px solid #eee', 'color': '#666', 'textAlign': 'center'}),
            html.Td(last_op, style={'padding': '8px', 'borderBottom': '1px solid #eee', 'color': '#666', 'textAlign': 'center'}),
            html.Td(
                "0 мин 00 сек",
                className="idle-timer",
                **{"data-start-sec": str(idle_sec), "data-render-time": str(now_ms), "data-on-shift": str(is_on_shift).lower()},
                style={'padding': '8px', 'borderBottom': '1px solid #eee', 'textAlign': 'center', 'color': '#666'}
            )
        ]))

    def icon_style(col):
        is_active = new_col == col and new_dir is not None
        return {'marginLeft': '5px', 'fontSize': '12px', 'opacity': '1', 'fontWeight': 'bold'} if is_active else {'marginLeft': '5px', 'fontSize': '12px', 'opacity': '0.5'}

    return rows, {'column': new_col, 'direction': new_dir}, icon_style('status'), icon_style('time'), icon_style('last_time'), icon_style('idle')


# ============================================================================
# ОСНОВНОЙ CALLBACK ЗАГРУЗКИ ДАННЫХ С КЭШИРОВАНИЕМ
# ============================================================================
@callback(
    [Output("performance-data-cache", "data", allow_duplicate=True),
     Output("fines-data", "data", allow_duplicate=True),
     Output("productivity-data", "data", allow_duplicate=True),
     Output("client-cache-store", "data"),
     Output("client-cache-timestamp", "data"),
     Output("etl-status-store", "data"),
     Output("etl-polling-interval", "disabled"),
     Output("data-refresh-trigger", "children")],
    [Input("refresh-button", "n_clicks"),
     Input("global-date-range", "data"),
     Input("etl-polling-interval", "n_intervals")],
    [State("client-cache-store", "data"),
     State("etl-status-store", "data")],
    prevent_initial_call='initial_duplicate'
)
def load_dashboard_data(refresh_clicks, date_range, n_intervals, client_cache, stored_etl_status):
    """
    Загрузка данных дашборда с учетом клиентского кэша и статуса ETL
    """
    ctx = dash.callback_context
    if not ctx.triggered:
        raise dash.exceptions.PreventUpdate
    
    trigger_id = ctx.triggered[0]['prop_id']
    logger.info(f"Load dashboard data triggered by: {trigger_id}")
    
    # ✅ НОВАЯ ПРОВЕРКА: инициализация даты если она пустая
    if not date_range or not date_range.get('start_date'):
        today = datetime.now().strftime('%Y-%m-%d')
        date_range = {'start_date': today, 'end_date': today}
        logger.info(f"Initialized date_range to today: {today}")
        
    start_date = date_range.get('start_date')
    end_date = date_range.get('end_date')
    
    # Проверяем статус ETL
    etl_status = ETLChecker.get_status()
    etl_is_updating = etl_status['is_updating']
    
    # Создаем ключ кэша
    cache_params = {
        'start_date': start_date,
        'end_date': end_date
    }
    cache_key = ClientCacheManager.create_cache_key('dashboard_data', cache_params)
    
    # Проверяем, нужно ли использовать клиентский кэш
    should_use_cache, cached_data = ClientCacheManager.should_use_cache(client_cache, etl_is_updating)
    
    # Если кэш валиден и его ключ совпадает - используем его
    if should_use_cache and cached_data and cached_data.get('_cache_key') == cache_key:
        logger.info("Using valid client cache")
        cached_performance = cached_data.get('data', {}).get('performance', [])
        cached_fines = cached_data.get('data', {}).get('fines', {})
        cached_productivity = cached_data.get('data', {}).get('productivity', [])
        
        # Если ETL обновляется - включаем polling
        polling_disabled = not etl_is_updating
        
        return [
            cached_performance,
            cached_fines,
            cached_productivity,
            client_cache,
            dash.no_update,
            etl_status,
            polling_disabled,
            datetime.now().isoformat()
        ]
    
    # Если ETL обновляется и нет валидного кэша - не загружаем данные
    if etl_is_updating:
        logger.info("ETL is updating, no valid cache - waiting...")
        return [
            dash.no_update,
            dash.no_update,
            dash.no_update,
            dash.no_update,
            dash.no_update,
            etl_status,
            False,
            dash.no_update
        ]
    
    # Загружаем свежие данные
    try:
        logger.info("Loading fresh data from database...")
        
        fresh_performance = get_performance_data(start_date, end_date)
        fresh_fines = get_fines_data(start_date, end_date)
        fresh_productivity = fresh_performance
        
        cache_package = {
            'performance': fresh_performance,
            'fines': fresh_fines,
            'productivity': fresh_productivity
        }
        
        cache_data = ClientCacheManager.prepare_cache_data(
            cache_package,
            ttl_seconds=300,
            is_stale=False
        )
        cache_data['_cache_key'] = cache_key
        
        client_cache_json = json.dumps(cache_data)
        logger.info("Fresh data loaded successfully")
        
        return [
            fresh_performance,
            fresh_fines,
            fresh_productivity,
            client_cache_json,
            datetime.now().isoformat(),
            etl_status,
            True,
            datetime.now().isoformat()
        ]
        
    except dash.exceptions.PreventUpdate:
        raise
    except Exception as e:
        logger.error(f"Error loading fresh data: {e}")
        if client_cache:
            try:
                cached = json.loads(client_cache) if isinstance(client_cache, str) else client_cache
                cached_performance = cached.get('data', {}).get('performance', [])
                cached_fines = cached.get('data', {}).get('fines', {})
                cached_productivity = cached.get('data', {}).get('productivity', [])
                return [
                    cached_performance, cached_fines, cached_productivity, client_cache, dash.no_update, etl_status, True, dash.no_update
                ]
            except: pass
        raise dash.exceptions.PreventUpdate

# ✅ ИЗМЕНЕНИЕ: Обновлённый колбэк с 3 выходами
@callback(
    Output('global-date-range', 'data', allow_duplicate=True),
    Output('global-date-start', 'value', allow_duplicate=True),
    Output('global-date-end', 'value', allow_duplicate=True),
    [Input('date-check-interval', 'n_intervals')],
    prevent_initial_call='initial_duplicate'
)
def update_global_date_range(n_intervals):
    """
    Обновление глобальной даты: синхронизация Store + видимых инпутов.
    Срабатывает каждые 5 минут (интервал date-check-interval).
    """
    today = datetime.now().strftime('%Y-%m-%d')
    date_data = {'start_date': today, 'end_date': today}
    return date_data, today, today


@callback(
    [Output('total-revisions-kpi', 'children'),
     Output('open-revisions-kpi', 'children'),
     Output('in-process-revisions-kpi', 'children'),
     Output('placement-errors-kpi', 'children'),
     Output('placement-correct-kpi', 'children'),
     Output('placement-errors-count-kpi', 'children'),
     Output('storage-cells-kpi', 'children'),
     Output('storage-cells-detail', 'children'),
     Output('order-accuracy-kpi', 'children')],
    [Input('global-date-range', 'data')],
    prevent_initial_call=False
)
def update_main_kpi_cards(date_range):
    """Обновление KPI карточек на главной вкладке"""
    try:
        revision_stats = get_revision_stats()
        storage_stats = get_storage_cells_stats()

        total_revisions = f"{revision_stats.get('total_revisions') or 0:,}"
        open_revisions = f"{revision_stats.get('open_revisions') or 0:,}"
        in_process_revisions = f"{revision_stats.get('in_process_revisions') or 0:,}"

        occupied = storage_stats.get('occupied_cells') or 0
        free = storage_stats.get('free_cells') or 0
        storage_kpi = f"{occupied}/{free}"
        
        occ_pct = storage_stats.get('occupied_percent') or 0
        free_pct = storage_stats.get('free_percent') or 0
        storage_detail = f"{occ_pct}% занято | {free_pct}% своб."
    except Exception as e:
        logger.error(f"Error getting revision/storage stats: {e}")
        total_revisions = "0"
        open_revisions = "0"
        in_process_revisions = "0"
        storage_kpi = "0/0"
        storage_detail = "0% занято | 0% своб."

    if not date_range:
        return (total_revisions, open_revisions, in_process_revisions, "0%", "0", "0", storage_kpi, storage_detail, "100%")

    start_date = date_range['start_date']
    end_date = date_range['end_date']

    try:
        placement_stats = get_placement_errors(start_date, end_date)
        accuracy, orders_without_errors, total_orders_accuracy, error_orders = get_order_accuracy(start_date, end_date)

        error_pct = placement_stats.get('error_percentage') or 0
        error_percentage = f"{error_pct}%"
        correct_count = f"{placement_stats.get('correct_count') or 0:,}"
        error_count = f"{placement_stats.get('error_count') or 0:,}"

        accuracy_val = accuracy if accuracy is not None else 0
        accuracy_str = f"{accuracy_val:.1f}%"

        return (total_revisions, open_revisions, in_process_revisions, error_percentage, correct_count, error_count, storage_kpi, storage_detail, accuracy_str)
    except Exception as e:
        logger.error(f"Error in update_main_kpi_cards: {e}")
        return ("0", "0", "0", "0%", "0", "0", "0/0", "0% занято | 0% своб.", "100%")


@callback(
    Output('last-update-time', 'children'),
    [Input('global-date-range', 'data'),
     Input('data-refresh-trigger', 'children')]
)
def update_last_update_time(date_range, trigger):
    return f"Обновлено: {datetime.now().strftime('%d.%m.%Y %H:%M')}"


@lru_cache(maxsize=32)
def _get_cached_data_for_dates(start_str, end_str):
    """Кэширует все данные для заданной пары дат, чтобы избежать повторных запросов к БД."""
    performance_data_cache = get_performance_data(start_str, end_str)
    shift_comparison_cache = get_shift_comparison(start_str, end_str)
    problematic_hours_cache = get_problematic_hours(start_str, end_str)
    error_hours_cache = get_error_hours_top_data(start_str, end_str)
    return performance_data_cache, shift_comparison_cache, problematic_hours_cache, error_hours_cache


@callback(
    [Output('global-date-range', 'data', allow_duplicate=True),
     Output('performance-data-cache', 'data', allow_duplicate=True),
     Output('shift-comparison-cache', 'data'),
     Output('problematic-hours-cache', 'data'),
     Output('error-hours-cache', 'data')],
    [Input('global-date-start', 'value'),
     Input('global-date-end', 'value')],
    prevent_initial_call='initial_duplicate'
)
def update_global_date_range_and_data(start_date, end_date):
    """Автоматическое обновление данных при изменении дат в полях ввода"""
    if not start_date or not end_date:
        raise dash.exceptions.PreventUpdate

    if isinstance(start_date, str):
        start_str = start_date[:10]
    else:
        start_str = start_date.strftime('%Y-%m-%d') if hasattr(start_date, 'strftime') else str(start_date)[:10]

    if isinstance(end_date, str):
        end_str = end_date[:10]
    else:
        end_str = end_date.strftime('%Y-%m-%d') if hasattr(end_date, 'strftime') else str(end_date)[:10]

    try:
        etl_status = ETLChecker.get_status()
        if etl_status['is_updating']:
            logger.info("ETL is updating, skipping data refresh - using cached data")
            raise dash.exceptions.PreventUpdate
            
        refresh_data(start_str, end_str)
        performance_data_cache, shift_comparison_cache, problematic_hours_cache, error_hours_cache = \
            _get_cached_data_for_dates(start_str, end_str)

        return {
            'start_date': start_str,
            'end_date': end_str
        }, performance_data_cache, shift_comparison_cache, problematic_hours_cache, error_hours_cache
        
    except dash.exceptions.PreventUpdate:
        raise
    except Exception as e:
        logger.error(f"Error in update_global_date_range_and_data: {e}")
        raise dash.exceptions.PreventUpdate


@callback(
    [Output("storage-cells-modal", "className"),
     Output("storage-modal-content", "className")],
    [Input("open-storage-modal", "n_clicks"),
     Input("close-storage-modal", "n_clicks")],
    [State("storage-cells-modal", "className"),
     State("storage-modal-content", "className")],
    prevent_initial_call=True
)
def toggle_storage_modal(open_clicks, close_clicks, modal_class, content_class):
    ctx = dash.callback_context
    if not ctx.triggered:
        return modal_class, content_class
    
    button_id = ctx.triggered[0]['prop_id'].split('.')[0]
    if button_id == 'open-storage-modal':
        return 'modal-visible', 'modal-content-visible'
    elif button_id == 'close-storage-modal':
        return 'modal-hidden', 'modal-content'
    return modal_class, content_class


@callback(
    [Output("revision-detail-modal", "className"),
     Output("revision-detail-modal-content", "className"),
     Output("revision-detail-table-body", "children"),
     Output("revision-detail-data-cache", "data")],
    [Input("open-revision-info", "n_clicks"),
     Input("close-revision-detail-modal", "n_clicks")],
    [State("revision-detail-modal", "className"),
     State("revision-detail-modal-content", "className")],
    prevent_initial_call=True
)
def toggle_revision_detail_modal(open_clicks, close_clicks, modal_class, content_class):
    ctx = dash.callback_context
    if not ctx.triggered:
        return modal_class, content_class, [], []
    
    button_id = ctx.triggered[0]['prop_id'].split('.')[0]
    if button_id == 'close-revision-detail-modal':
        return 'modal-hidden', 'modal-content', [], []
    
    if button_id == 'open-revision-info' and open_clicks:
        from data.queries_mssql import get_revision_detail_data
        detail_data = get_revision_detail_data()
        
        table_rows = []
        for item in detail_data:
            status_color = '#4CAF50' if item['status_rus'] == 'Открыто' else '#FF9800'
            variance_color = '#4CAF50' if item['variance'] == 0 else '#F44336'
            table_rows.append(
                html.Tr([
                    html.Td(str(item['internal_count_num']), style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px', 'fontWeight': 'bold'}),
                    html.Td(html.Span(item['status_rus'], style={'padding': '4px 8px', 'borderRadius': '4px', 'fontSize': '11px', 'fontWeight': 'bold', 'color': 'white', 'backgroundColor': status_color}), style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px'}),
                    html.Td(item['item'], style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px'}),
                    html.Td(item['item_desc'] or '—', style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px', 'color': '#555', 'maxWidth': '200px', 'overflow': 'hidden', 'textOverflow': 'ellipsis'}),
                    html.Td(item['lot'], style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px'}),
                    html.Td(item['location'], style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px'}),
                    html.Td(f"{item['quantity_counted']:,.2f}", style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px', 'textAlign': 'right'}),
                    html.Td(f"{item['system_quantity']:,.2f}", style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px', 'textAlign': 'right'}),
                    html.Td(html.Span(f"{item['variance']:+,.2f}", style={'color': variance_color, 'fontWeight': 'bold'}), style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px', 'textAlign': 'right'}),
                    html.Td(item['counted_by_user'] or '—', style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px', 'color': '#555'}),
                    html.Td(item['counted_date_time'] or '—', style={'padding': '8px', 'borderBottom': '1px solid #eee', 'fontSize': '12px', 'color': '#555'})
                ])
            )
        if not table_rows:
            table_rows = [html.Tr([html.Td("Нет данных", colSpan=11, style={'textAlign': 'center', 'padding': '20px', 'color': '#666'})])]
        return 'modal-visible', 'modal-content-visible', table_rows, detail_data
    
    return modal_class, content_class, [], []


# Callback для печати ревизий — сохраняет данные и возвращает URL
@callback(
    Output("print-revision-url-store", "data"),
    Input("print-revision-btn", "n_clicks"),
    State("revision-detail-data-cache", "data"),
    prevent_initial_call=True
)
def print_revision_prepare_url(n_clicks, cached_data):
    """Сохраняет данные во временный файл и возвращает URL"""
    if not cached_data:
        raise dash.exceptions.PreventUpdate
    
    import json
    import os
    import uuid
    import tempfile
    import time
    from data.queries_mssql import get_items_locations
    
    item_codes = [item.get('item', '') for item in cached_data if item.get('item')]
    locations_map = get_items_locations(item_codes) if item_codes else {}
    
    key = uuid.uuid4().hex
    
    temp_dir = os.path.join(tempfile.gettempdir(), 'dash_print')
    os.makedirs(temp_dir, exist_ok=True)
    
    # Удаляем файлы старше 1 часа
    try:
        now = time.time()
        for filename in os.listdir(temp_dir):
            filepath = os.path.join(temp_dir, filename)
            if os.path.isfile(filepath) and filename.endswith('.json'):
                if now - os.path.getmtime(filepath) > 3600:
                    os.remove(filepath)
    except:
        pass
    
    data_file = os.path.join(temp_dir, f'{key}.json')
    with open(data_file, 'w', encoding='utf-8') as f:
        json.dump({
            'revisions': cached_data,
            'locations': locations_map
        }, f, ensure_ascii=False)
    
    # Добавляем timestamp чтобы URL всегда был уникальным
    return f"/print-revision?key={key}&t={int(time.time())}"

# Callback для экспорта ревизий в Excel
@callback(
    Output("download-revision-csv", "data"),
    Input("export-revision-csv-btn", "n_clicks"),
    State("revision-detail-data-cache", "data"),
    prevent_initial_call=True
)
def export_revision_to_csv(n_clicks, cached_data):
    """Экспорт данных ревизий в Excel с остатками по всем ячейкам товара"""
    if not cached_data:
        return None
    
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
    from data.queries_mssql import get_items_locations
    
    item_codes = [item['item'] for item in cached_data if item.get('item')]
    locations_map = get_items_locations(item_codes)
    
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ревизии'
    
    item_font = Font(name='Arial', size=11, bold=True, color='1976D2')
    revision_font = Font(name='Arial', size=10, color='333333')
    subheader_font = Font(name='Arial', size=9, bold=True, color='666666')
    location_font = Font(name='Arial', size=9, color='555555')
    separator_font = Font(name='Arial', size=8, color='CCCCCC')
    
    col_widths = {'A': 18, 'B': 16, 'C': 14, 'D': 10, 'E': 18, 'F': 14, 'G': 18, 'H': 20}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    row = 1
    
    for item_data in cached_data:
        item_code = item_data.get('item', '')
        item_desc = item_data.get('item_desc', '')
        
        max_desc_len = 80
        if item_desc and len(item_desc) > max_desc_len:
            short_desc = item_desc[:max_desc_len - 3] + '...'
        else:
            short_desc = item_desc or 'Без описания'
        
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=f"ТОВАР: {item_code} — {short_desc}")
        cell.font = item_font
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        row += 1
        
        revision_line_1 = (
            f"Ревизия №{item_data['internal_count_num']} | "
            f"Статус: {item_data['status_rus']} | "
            f"Локация: {item_data.get('location', '—')} | "
            f"Подсчитано: {item_data['quantity_counted']:.2f} | "
            f"Системное: {item_data['system_quantity']:.2f} | "
            f"Отклонение: {item_data['variance']:+.2f}"
        )
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=revision_line_1)
        cell.font = revision_font
        row += 1
        
        revision_line_2 = (
            f"Выполнил: {item_data.get('counted_by_user', '—')} | "
            f"Дата: {item_data.get('counted_date_time', '—')}"
        )
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=revision_line_2)
        cell.font = revision_font
        row += 1
        
        other_locations = locations_map.get(str(item_code), [])
        
        if other_locations:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            cell = ws.cell(row=row, column=1, value='ОСТАТКИ НА ВСЕХ ЯЧЕЙКАХ:')
            cell.font = subheader_font
            row += 1
            
            loc_headers = ['Локация', 'Количество', 'Ед. изм.', 'Партия']
            for col_idx, header in enumerate(loc_headers):
                cell = ws.cell(row=row, column=col_idx + 1, value=header)
                cell.font = Font(name='Arial', size=9, bold=True, color='FFFFFF')
                cell.fill = PatternFill(start_color='999999', end_color='999999', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
                cell.border = Border(
                    left=Side(style='thin', color='CCCCCC'),
                    right=Side(style='thin', color='CCCCCC'),
                    top=Side(style='thin', color='CCCCCC'),
                    bottom=Side(style='thin', color='CCCCCC')
                )
            row += 1
            
            for loc in other_locations:
                cell_a = ws.cell(row=row, column=1, value=loc['location'])
                cell_a.font = location_font
                cell_a.alignment = Alignment(horizontal='left')
                cell_a.border = Border(left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'), bottom=Side(style='thin', color='E0E0E0'))
                
                cell_b = ws.cell(row=row, column=2, value=loc['quantity'])
                cell_b.font = location_font
                cell_b.alignment = Alignment(horizontal='right')
                cell_b.number_format = '#,##0'
                cell_b.border = Border(left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'), bottom=Side(style='thin', color='E0E0E0'))
                
                cell_c = ws.cell(row=row, column=3, value=loc.get('um', 'шт'))
                cell_c.font = location_font
                cell_c.alignment = Alignment(horizontal='center')
                cell_c.border = Border(left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'), bottom=Side(style='thin', color='E0E0E0'))
                
                cell_d = ws.cell(row=row, column=4, value=loc.get('lot', '') if loc.get('lot') else '—')
                cell_d.font = location_font
                cell_d.alignment = Alignment(horizontal='left')
                cell_d.border = Border(left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'), bottom=Side(style='thin', color='E0E0E0'))
                row += 1
        else:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            cell = ws.cell(row=row, column=1, value='Нет данных об остатках на других ячейках')
            cell.font = Font(name='Arial', size=9, italic=True, color='999999')
            row += 1
        
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value='─' * 100)
        cell.font = separator_font
        row += 1
    
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.print_area = f'A1:H{row}'
    ws.sheet_properties.pageSetUpPr = ws.sheet_properties.pageSetUpPr or type(ws.sheet_properties.pageSetUpPr)()
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.scale = 85
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return dcc.send_bytes(buffer.getvalue(), "revision_detail.xlsx")

@callback(
    Output("storage-all-data", "data"),
    Input("storage-cells-modal", "className")
)
def load_storage_data(modal_class):
    if modal_class == 'modal-visible':
        all_data = get_all_storage_data()
        return {'all_data': all_data}
    return dash.no_update


@callback(
    [Output("filter-storage-type", "options"),
     Output("filter-locating-zone", "options"),
     Output("filter-allocation-zone", "options"),
     Output("filter-location-type", "options"),
     Output("filter-work-zone", "options"),
     Output("storage-total-cells", "children"),
     Output("storage-occupied-cells", "children"),
     Output("storage-free-cells", "children"),
     Output("storage-occupied-percent", "children"),
     Output("storage-empty-chart", "option"),
     Output("storage-types-pie-chart", "option"),
     Output("storage-types-bar-chart", "option"),
     Output("storage-current-filters", "data")],
    [Input("filter-storage-type", "value"),
     Input("filter-locating-zone", "value"),
     Input("filter-allocation-zone", "value"),
     Input("filter-location-type", "value"),
     Input("filter-work-zone", "value"),
     Input("filter-only-empty", "value"),
     Input("storage-all-data", "data")]
)
def update_storage_filters_and_charts(storage_type_val, locating_zone_val, allocation_zone_val, 
                                     location_type_val, work_zone_val, only_empty_val, all_data):
    if not all_data or 'all_data' not in all_data:
        empty_options = [{'label': 'Все', 'value': 'Все'}]
        empty_chart = {"title": {"text": "Нет данных", "left": "center"}}
        return (empty_options, empty_options, empty_options, empty_options, empty_options,
                "0", "0", "0", "0%", empty_chart, empty_chart, empty_chart,
                {'storage_type': 'Все', 'locating_zone': 'Все', 'allocation_zone': 'Все', 
                 'location_type': 'Все', 'work_zone': 'Все', 'only_empty': False})
    
    try:
        is_only_empty = True if only_empty_val and 'empty' in only_empty_val else False
        current_filters = {
            'storage_type': storage_type_val if storage_type_val != 'Все' else None,
            'locating_zone': locating_zone_val if locating_zone_val != 'Все' else None,
            'allocation_zone': allocation_zone_val if allocation_zone_val != 'Все' else None,
            'location_type': location_type_val if location_type_val != 'Все' else None,
            'work_zone': work_zone_val if work_zone_val != 'Все' else None,
            'only_empty': is_only_empty
        }
        
        filtered_result = filter_storage_data(all_data['all_data'], current_filters)
        available_filters = filtered_result['available_filters']
        
        def create_options(values):
            options = [{'label': 'Все', 'value': 'Все'}]
            for value in values:
                if value and value.strip():
                    options.append({'label': value, 'value': value})
            return options
        
        storage_type_options = create_options(available_filters['storage_type'])
        locating_zone_options = create_options(available_filters['locating_zone'])
        allocation_zone_options = create_options(available_filters['allocation_zone'])
        location_type_options = create_options(available_filters['location_type'])
        work_zone_options = create_options(available_filters['work_zone'])
        
        summary = filtered_result['summary']
        if is_only_empty:
            total_cells = f"{summary['empty']:,}"
            occupied_cells = "0"
            free_cells = f"{summary['empty']:,}"
            occupied_percent = 0
        else:
            total_cells = f"{summary['total']:,}"
            occupied_cells = f"{summary['occupied']:,}"
            free_cells = f"{summary['empty']:,}"
            occupied_percent = round((summary['occupied'] / summary['total']) * 100, 1) if summary['total'] > 0 else 0
            
        occupied_percent_str = f"{occupied_percent}%"
        
        empty_chart = create_empty_pie_chart(summary, current_filters)
        types_pie_chart = create_types_pie_chart(filtered_result['chart_data'], current_filters)
        types_bar_chart = create_types_bar_chart(filtered_result['chart_data'], current_filters)
        
        current_filter_values = {
            'storage_type': storage_type_val or 'Все', 'locating_zone': locating_zone_val or 'Все',
            'allocation_zone': allocation_zone_val or 'Все', 'location_type': location_type_val or 'Все',
            'work_zone': work_zone_val or 'Все', 'only_empty': is_only_empty
        }
        
        return (storage_type_options, locating_zone_options, allocation_zone_options, location_type_options,
                work_zone_options, total_cells, occupied_cells, free_cells, occupied_percent_str,
                empty_chart, types_pie_chart, types_bar_chart, current_filter_values)
    except Exception as e:
        logger.error(f"ERROR in callback: {e}")
        empty_options = [{'label': 'Все', 'value': 'Все'}]
        empty_chart = {"title": {"text": "Ошибка", "left": "center"}}
        return (empty_options, empty_options, empty_options, empty_options, empty_options,
                "0", "0", "0", "0%", empty_chart, empty_chart, empty_chart,
                {'storage_type': 'Все', 'locating_zone': 'Все', 'allocation_zone': 'Все', 
                 'location_type': 'Все', 'work_zone': 'Все', 'only_empty': False})


@callback(
    [Output("filter-storage-type", "value"),
     Output("filter-locating-zone", "value"),
     Output("filter-allocation-zone", "value"),
     Output("filter-location-type", "value"),
     Output("filter-work-zone", "value")],
    [Input("reset-all-filters-btn", "n_clicks"),
     Input("reset-storage-type-btn", "n_clicks"),
     Input("reset-locating-zone-btn", "n_clicks"),
     Input("reset-allocation-zone-btn", "n_clicks"),
     Input("reset-location-type-btn", "n_clicks"),
     Input("reset-work-zone-btn", "n_clicks")],
    [State("filter-storage-type", "value"),
     State("filter-locating-zone", "value"),
     State("filter-allocation-zone", "value"),
     State("filter-location-type", "value"),
     State("filter-work-zone", "value")],
    prevent_initial_call=True
)
def reset_filters(all_clicks, storage_clicks, locating_clicks, allocation_clicks, location_clicks, work_clicks,
                  current_storage, current_locating, current_allocation, current_location, current_work):
    ctx = dash.callback_context
    if not ctx.triggered:
        raise dash.exceptions.PreventUpdate
    
    button_id = ctx.triggered[0]['prop_id'].split('.')[0]
    if button_id == "reset-all-filters-btn":
        return ['Все', 'Все', 'Все', 'Все', 'Все']
    
    new_storage = current_storage
    new_locating = current_locating
    new_allocation = current_allocation
    new_location = current_location
    new_work = current_work
    
    if button_id == "reset-storage-type-btn": new_storage = 'Все'
    elif button_id == "reset-locating-zone-btn": new_locating = 'Все'
    elif button_id == "reset-allocation-zone-btn": new_allocation = 'Все'
    elif button_id == "reset-location-type-btn": new_location = 'Все'
    elif button_id == "reset-work-zone-btn": new_work = 'Все'
    
    return new_storage, new_locating, new_allocation, new_location, new_work


# ============================================================================
# КОЛБЭКИ ТАБЛИЦЫ СОТРУДНИКОВ
# ============================================================================

@callback(
    [Output('shift-employees-table-body', 'children', allow_duplicate=True),
     Output('shift-table-sort-state', 'data', allow_duplicate=True),
     Output('sort-status-icon', 'style', allow_duplicate=True),
     Output('sort-time-icon', 'style', allow_duplicate=True),
     Output('sort-last-time-icon', 'style', allow_duplicate=True),
     Output('sort-idle-icon', 'style', allow_duplicate=True)],
    [Input('shift-table-interval', 'n_intervals'),
     Input('position-filter', 'value'),
     Input('brigade-filter', 'value')],
    prevent_initial_call='initial_duplicate'
)
def update_shift_employees_table(n_intervals, position_filter, brigade_filter):
    # Принудительно сбрасываем на ASC при каждом обновлении таймера или фильтров
    return _apply_sorting('time', {'column': 'time', 'force_dir': 'asc'}, position_filter, brigade_filter)


# 2. Сортировка по статусу
@callback(
    [Output('shift-employees-table-body', 'children', allow_duplicate=True),
     Output('shift-table-sort-state', 'data', allow_duplicate=True),
     Output('sort-status-icon', 'style', allow_duplicate=True),
     Output('sort-time-icon', 'style', allow_duplicate=True),
     Output('sort-last-time-icon', 'style', allow_duplicate=True),
     Output('sort-idle-icon', 'style', allow_duplicate=True)],
    [Input('sort-status-header', 'n_clicks')],
    [State('shift-table-sort-state', 'data'), State('position-filter', 'value'), State('brigade-filter', 'value')],
    prevent_initial_call=True
)
def sort_by_status(n_clicks, sort_state, position_filter, brigade_filter):
    if not n_clicks: raise dash.exceptions.PreventUpdate
    return _apply_sorting('status', sort_state, position_filter, brigade_filter)


# 3. Сортировка по времени первой операции
@callback(
    [Output('shift-employees-table-body', 'children', allow_duplicate=True),
     Output('shift-table-sort-state', 'data', allow_duplicate=True),
     Output('sort-status-icon', 'style', allow_duplicate=True),
     Output('sort-time-icon', 'style', allow_duplicate=True),
     Output('sort-last-time-icon', 'style', allow_duplicate=True),
     Output('sort-idle-icon', 'style', allow_duplicate=True)],
    [Input('sort-time-header', 'n_clicks')],
    [State('shift-table-sort-state', 'data'), State('position-filter', 'value'), State('brigade-filter', 'value')],
    prevent_initial_call=True
)
def sort_by_time(n_clicks, sort_state, position_filter, brigade_filter):
    if not n_clicks: raise dash.exceptions.PreventUpdate
    return _apply_sorting('time', sort_state, position_filter, brigade_filter)


# 4. Сортировка по времени последней операции
@callback(
    [Output('shift-employees-table-body', 'children', allow_duplicate=True),
     Output('shift-table-sort-state', 'data', allow_duplicate=True),
     Output('sort-status-icon', 'style', allow_duplicate=True),
     Output('sort-time-icon', 'style', allow_duplicate=True),
     Output('sort-last-time-icon', 'style', allow_duplicate=True),
     Output('sort-idle-icon', 'style', allow_duplicate=True)],
    [Input('sort-last-time-header', 'n_clicks')],
    [State('shift-table-sort-state', 'data'), State('position-filter', 'value'), State('brigade-filter', 'value')],
    prevent_initial_call=True
)
def sort_by_last_time(n_clicks, sort_state, position_filter, brigade_filter):
    if not n_clicks: raise dash.exceptions.PreventUpdate
    return _apply_sorting('last_time', sort_state, position_filter, brigade_filter)


# 5. Сортировка по времени простоя
@callback(
    [Output('shift-employees-table-body', 'children', allow_duplicate=True),
     Output('shift-table-sort-state', 'data', allow_duplicate=True),
     Output('sort-status-icon', 'style', allow_duplicate=True),
     Output('sort-time-icon', 'style', allow_duplicate=True),
     Output('sort-last-time-icon', 'style', allow_duplicate=True),
     Output('sort-idle-icon', 'style', allow_duplicate=True)],
    [Input('sort-idle-header', 'n_clicks')],
    [State('shift-table-sort-state', 'data'), State('position-filter', 'value'), State('brigade-filter', 'value')],
    prevent_initial_call=True
)
def sort_by_idle(n_clicks, sort_state, position_filter, brigade_filter):
    if not n_clicks: raise dash.exceptions.PreventUpdate
    return _apply_sorting('idle', sort_state, position_filter, brigade_filter)


# ============================================================================
# CALLBACK ДЛЯ ТАБЛИЦЫ ПРОИЗВОДИТЕЛЬНОСТИ (С ПОЛНЫМИ СТИЛЯМИ)
# ============================================================================
@callback(
    [Output('productivity-table-body', 'children'),
     Output('prod-sort-state', 'data'),
     Output('prod-sort-Сотрудник-icon', 'children'),
     Output('prod-sort-Должность-icon', 'children'),
     Output('prod-sort-Операции-icon', 'children'),
     Output('prod-sort-Объем-icon', 'children'),
     Output('prod-sort-Контейнеров-icon', 'children'),
     Output('prod-sort-Ср_объем-icon', 'children'),
     Output('prod-sort-Время-icon', 'children'),
     Output('prod-sort-Оп/час-icon', 'children'),
     Output('prod-sort-Заработок-icon', 'children')],
    [Input('global-date-range', 'data'),
     Input('prod-sort-Сотрудник-header', 'n_clicks'),
     Input('prod-sort-Должность-header', 'n_clicks'),
     Input('prod-sort-Операции-header', 'n_clicks'),
     Input('prod-sort-Объем-header', 'n_clicks'),
     Input('prod-sort-Контейнеров-header', 'n_clicks'),
     Input('prod-sort-Ср_объем-header', 'n_clicks'),
     Input('prod-sort-Время-header', 'n_clicks'),
     Input('prod-sort-Оп/час-header', 'n_clicks'),
     Input('prod-sort-Заработок-header', 'n_clicks'),
     Input('date-check-interval', 'n_intervals')],
    [State('prod-sort-state', 'data')],
    prevent_initial_call=False
)
def update_productivity_table(date_range, emp_clicks, pos_clicks, ops_clicks, vol_clicks, cont_clicks, avg_vol_clicks, time_clicks, ops_hr_clicks, earn_clicks, n_intervals, sort_state):
    ctx = dash.callback_context
    if not date_range:
        raise dash.exceptions.PreventUpdate
    
    start_date = date_range['start_date']
    end_date = date_range['end_date']

    if sort_state is None:
        sort_state = {'column': 'Заработок', 'direction': 'desc'}

    if ctx.triggered:
        trigger_id = ctx.triggered[0]['prop_id'].split('.')[0]
        col_map = {
            'prod-sort-Сотрудник-header': 'Сотрудник',
            'prod-sort-Должность-header': 'Должность',
            'prod-sort-Операции-header': 'Операции',
            'prod-sort-Объем-header': 'Объем',
            'prod-sort-Контейнеров-header': 'Контейнеров',
            'prod-sort-Ср_объем-header': 'Ср_объем',
            'prod-sort-Время-header': 'Время',
            'prod-sort-Оп/час-header': 'Оп/час',
            'prod-sort-Заработок-header': 'Заработок'
        }
        if trigger_id in col_map:
            clicked_col = col_map[trigger_id]
            if sort_state['column'] == clicked_col:
                sort_state['direction'] = 'asc' if sort_state['direction'] == 'desc' else 'desc'
            else:
                sort_state['column'] = clicked_col
                sort_state['direction'] = 'desc'

    try:
        from data.mssql_client import clear_cache
        clear_cache()
    except:
        pass
    
    try:
        from data.queries_mssql import get_performance_data
        data = get_performance_data(start_date, end_date)
    except Exception as e:
        logger.error(f"Ошибка загрузки performance_data: {e}")
        data = []
        
    if not data:
        return [html.Tr([html.Td("Нет данных за выбранный период", colSpan=9, style={'textAlign': 'center', 'padding': '20px', 'color': '#666'})])], sort_state, '', '', '', '', '', '', '', '', ''

    df = pd.DataFrame(data)
    col_name = sort_state['column']
    ascending = sort_state['direction'] == 'asc'

    df_col_map = {
        'Сотрудник': 'Сотрудник',
        'Должность': 'position',
        'Операции': 'Общее_кол_операций',
        'Объем': 'Объем',
        'Контейнеров': 'Контейнеров',
        'Ср_объем': 'Ср_объем_контейнера',
        'Время': 'Ср_время_на_операцию',
        'Оп/час': 'Операций_в_час',
        'Заработок': 'Заработок'
    }
    
    if col_name in df_col_map:
        df = df.sort_values(by=df_col_map[col_name], ascending=ascending)

    icons = {}
    for key in ['Сотрудник', 'Должность', 'Операции', 'Объем', 'Контейнеров', 'Ср_объем', 'Время', 'Оп/час', 'Заработок']:
        if sort_state['column'] == key:
            icons[key] = ' ▲' if sort_state['direction'] == 'asc' else ' ▼'
        else:
            icons[key] = ''

    rows = []
    # ✅ ИСПРАВЛЕНИЕ: используем enumerate для стабильного индекса, 
    # но в id ссылки передаём user_name как уникальный идентификатор
    for row_idx, (idx, row) in enumerate(df.iterrows()):
        ops_per_hour = row.get('Операций_в_час', 0)
        earnings = row.get('Заработок', 0)

        emp_display_name = row.get('Сотрудник', 'Неизвестно')
        emp_position = row.get('position', 'Не указана')
        user_name = row.get('user_name', '')  # ✅ Уникальный идентификатор
        
        volume = row.get('Объем', 0)
        volume_str = f"{volume:.2f} м³" if volume > 0 else "-"

        containers = row.get('Контейнеров', 0)
        cont_str = str(containers) if containers > 0 else "-"

        avg_vol = row.get('Ср_объем_контейнера', 0)
        avg_vol_str = f"{avg_vol:.2f} м³" if avg_vol > 0 else "-"

        rows.append(html.Tr([
            html.Td(
                # ✅ КЛЮЧЕВОЕ ИСПРАВЛЕНИЕ: передаём user_name вместо числового индекса
                html.A(emp_display_name, href='#', id={'type': 'employee', 'index': user_name}, className='employee-link'), 
                style={'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'color': '#333'}
            ),
            html.Td(
                emp_position, 
                style={'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'color': '#333', 'textAlign': 'center'}
            ),
            html.Td(
                str(int(row.get('Общее_кол_операций', 0))), 
                style={'color': '#333', 'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'fontWeight': 'bold', 'textAlign': 'center'}
            ),
            html.Td(
                volume_str, 
                style={'color': '#333', 'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'fontWeight': 'bold', 'textAlign': 'center'}
            ),
            html.Td(
                cont_str, 
                style={'color': '#333', 'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'fontWeight': 'bold', 'textAlign': 'center'}
            ),
            html.Td(
                avg_vol_str, 
                style={'color': '#333', 'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'fontWeight': 'bold', 'textAlign': 'center'}
            ),
            html.Td(
                f"{row.get('Ср_время_на_операцию', 0):.1f} мин", 
                style={'color': '#333', 'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'fontWeight': 'bold', 'textAlign': 'center'}
            ),
            html.Td(
                f"{ops_per_hour:.1f}", 
                style={'color': '#333', 'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'fontWeight': 'bold', 'textAlign': 'center'}
            ),
            html.Td(
                f"{earnings:,.2f} ₽", 
                style={'color': '#333', 'padding': '12px', 'borderBottom': '1px solid #eee', 'fontSize': '14px', 'fontWeight': 'bold', 'textAlign': 'center'}
            )
        ]))

    return (
        rows, sort_state, 
        icons['Сотрудник'], icons['Должность'], icons['Операции'], 
        icons['Объем'], icons['Контейнеров'], icons['Ср_объем'], 
        icons['Время'], icons['Оп/час'], icons['Заработок']
    )


# Callback для обновления статистики смены на главной вкладке
@callback(
    Output('shift-stats-info', 'children'),
    [Input('global-date-range', 'data')]
)
def update_shift_stats_info(date_range):
    """Обновление информации о смене в общей сводке с данными мониторинга нагрузки групп"""
    try:
        groups_data = get_group_load_monitor()
        if not groups_data:
            return html.Div("Смена отдыхает или нет данных о нагрузке", style={'color': '#666', 'textAlign': 'center', 'padding': '20px', 'fontSize': '14px'})
        
        color_map = {'GREEN': '#4CAF50', 'YELLOW': '#FF9800', 'RED': '#F44336', 'GRAY': '#9E9E9E'}
        table_rows = []
        
        table_rows.append(html.Thead([html.Tr([
            html.Th('Группа', style={'color': '#666', 'padding': '12px', 'textAlign': 'left', 'fontSize': '14px', 'borderBottom': '2px solid #eee', 'background': '#f8f9fa'}),
            html.Th('Тип работы', style={'color': '#666', 'padding': '12px', 'textAlign': 'left', 'fontSize': '14px', 'borderBottom': '2px solid #eee', 'background': '#f8f9fa'}),
            html.Th('Кол-во сотрудников', style={'color': '#666', 'padding': '12px', 'textAlign': 'center', 'fontSize': '14px', 'borderBottom': '2px solid #eee', 'background': '#f8f9fa'}),
            html.Th('Кол-во заданий', style={'color': '#666', 'padding': '12px', 'textAlign': 'center', 'fontSize': '14px', 'borderBottom': '2px solid #eee', 'background': '#f8f9fa'}),
            html.Th('Статус', style={'color': '#666', 'padding': '12px', 'textAlign': 'center', 'fontSize': '14px', 'borderBottom': '2px solid #eee', 'background': '#f8f9fa'})
        ])]))
        
        tbody_rows = []
        for item in groups_data:
            group_name = item.get('group_name', '')
            work_type = item.get('work_type', '')
            employees_count = item.get('employees_count', 0)
            open_tasks_count = item.get('open_tasks_count', 0)
            status_color = item.get('status_color', 'GRAY')
            color = color_map.get(status_color, '#9E9E9E')
            
            tbody_rows.append(html.Tr([
                html.Td(group_name, style={'padding': '10px 12px', 'fontSize': '13px', 'borderBottom': '1px solid #f0f0f0'}),
                html.Td(work_type, style={'padding': '10px 12px', 'fontSize': '13px', 'borderBottom': '1px solid #f0f0f0'}),
                html.Td(str(employees_count), style={'padding': '10px 12px', 'fontSize': '13px', 'borderBottom': '1px solid #f0f0f0', 'textAlign': 'center', 'fontWeight': 'bold'}),
                html.Td(str(open_tasks_count), style={'padding': '10px 12px', 'fontSize': '13px', 'borderBottom': '1px solid #f0f0f0', 'textAlign': 'center', 'fontWeight': 'bold'}),
                html.Td(html.Span('●', style={'color': color, 'fontSize': '28px', 'fontWeight': 'bold'}), style={'padding': '10px 12px', 'textAlign': 'center', 'borderBottom': '1px solid #f0f0f0'})
            ], className='table-row-hover'))
        
        table_rows.append(html.Tbody(tbody_rows))
        stats_table = html.Table(table_rows, style={'width': '100%', 'borderCollapse': 'collapse'})
        
        return html.Div([
            html.H4("Мониторинг нагрузки групп", style={'marginBottom': '12px', 'color': '#1976d2', 'fontSize': '18px', 'fontWeight': 'bold'}),
            html.Div(stats_table, style={'maxHeight': '450px', 'overflowY': 'auto'})
        ], style={'height': '100%', 'overflow': 'hidden'})
    except Exception as e:
        logger.error(f"Error in update_shift_stats_info: {e}")
        return html.Div(f"Ошибка загрузки данных: {str(e)}", style={'color': '#F44336', 'padding': '15px', 'textAlign': 'center', 'fontSize': '14px'})